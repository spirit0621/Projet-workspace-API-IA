import os
import sys
from pathlib import Path

current_dir = Path(__file__).resolve().parent
src_dir = current_dir.parent / "src"
if str(src_dir) not in sys.path:
    sys.path.insert(0, str(src_dir))

docx_path = r"C:\Users\alves\Desktop\Lycée, bts , formation, master\CFA-insta\Master 1 SI\TP\TP2\TP_Fil_Rouge_Assistant_IA_Local_NB_Dev_REPONSES.docx"
md_source_path = r"C:\Users\alves\Desktop\Lycée, bts , formation, master\CFA-insta\Master 1 SI\TP\TP2\TP2_FIL_ROUGE_ASSISTANT_IA_LOCAL_REPONSES_COMPLETES.md"

def build_with_docx():
    import docx
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = docx.Document()

    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = title_p.add_run("TP FIL ROUGE : Assistant IA Métier Local\n")
    r_title.bold = True
    r_title.font.size = Pt(20)
    r_title.font.color.rgb = RGBColor(0, 51, 102)

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = sub_p.add_run("Document de Compte-Rendu Intégral avec Réponses sous Chaque Question\n")
    r_sub.bold = True
    r_sub.font.size = Pt(13)
    r_sub.font.color.rgb = RGBColor(100, 100, 100)

    meta_p = doc.add_paragraph()
    meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_meta = meta_p.add_run("CFA-INSTA • Master 1 Architecture des SI • Module IA & MCP • TECHCORP\nProfil : Développement Applicatif (DevAssist-GPT)\n")
    r_meta.italic = True
    r_meta.font.size = Pt(10.5)

    doc.add_paragraph("―" * 45).alignment = WD_ALIGN_PARAGRAPH.CENTER

    if not os.path.exists(md_source_path):
        print(f"Fichier source markdown introuvable : {md_source_path}")
        return

    with open(md_source_path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    in_code = False
    code_block = []

    for line in lines:
        raw_line = line.rstrip("\n")
        
        if raw_line.startswith("```"):
            if in_code:
                in_code = False
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.3)
                r = p.add_run("\n".join(code_block))
                r.font.name = "Consolas"
                r.font.size = Pt(9)
                r.font.color.rgb = RGBColor(40, 40, 40)
                code_block = []
            else:
                in_code = True
                code_block = []
            continue

        if in_code:
            code_block.append(raw_line)
            continue

        stripped = raw_line.strip()
        if not stripped:
            continue

        if stripped.startswith("# "):
            p = doc.add_heading(level=1)
            r = p.add_run(stripped[2:])
            r.font.size = Pt(16)
            r.bold = True
            r.font.color.rgb = RGBColor(0, 51, 102)
        elif stripped.startswith("## "):
            p = doc.add_heading(level=2)
            r = p.add_run(stripped[3:])
            r.font.size = Pt(13)
            r.bold = True
            r.font.color.rgb = RGBColor(30, 80, 140)
        elif stripped.startswith("### "):
            p = doc.add_heading(level=3)
            r = p.add_run(stripped[4:])
            r.font.size = Pt(11.5)
            r.bold = True
            r.font.color.rgb = RGBColor(60, 60, 60)
        elif stripped.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.2)
            r = p.add_run(stripped[2:])
            r.italic = True
            r.bold = True
            r.font.color.rgb = RGBColor(0, 70, 150)
        elif stripped.startswith("- [x]") or stripped.startswith("- [ ]"):
            p = doc.add_paragraph(style="List Bullet")
            symbol = "☑ " if "[x]" in stripped else "☐ "
            r = p.add_run(symbol + stripped[5:].strip())
            r.font.size = Pt(10)
        elif stripped.startswith("- ") or stripped.startswith("• "):
            p = doc.add_paragraph(style="List Bullet")
            r = p.add_run(stripped[2:])
            r.font.size = Pt(10)
        elif stripped.startswith("|"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.1)
            r = p.add_run(stripped)
            r.font.name = "Consolas"
            r.font.size = Pt(8.5)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(stripped)
            r.font.size = Pt(10.5)

    doc.save(docx_path)
    print(f"✅ Document Word généré avec succès avec python-docx : {docx_path}")

def build_with_google_docs():
    print("Tentative de génération via Google Docs & Drive API...")
    try:
        import auth
        from services import assistant_docs
    except ImportError:
        from src import auth
        from src.services import assistant_docs

    creds = auth.get_credentials()
    docs_service = assistant_docs.get_service(creds)

    if not os.path.exists(md_source_path):
        print(f"Fichier source markdown introuvable : {md_source_path}")
        return

    with open(md_source_path, "r", encoding="utf-8") as f:
        full_text = f.read()

    doc_title = "TP FIL ROUGE - Assistant IA Métier Local (DevAssist-GPT)"
    doc = docs_service.documents().create(body={'title': doc_title}).execute()
    doc_id = doc.get('documentId')

    docs_service.documents().batchUpdate(
        documentId=doc_id,
        body={'requests': [{'insertText': {'location': {'index': 1}, 'text': full_text}}]}
    ).execute()

    assistant_docs.export_word(creds, doc_id, docx_path)
    print(f"✅ Document Word généré avec succès via Google Docs API : {docx_path}")

if __name__ == "__main__":
    try:
        build_with_docx()
    except ImportError:
        print("python-docx non installé, bascule vers Google Docs API...")
        build_with_google_docs()
    except Exception as e:
        print(f"Erreur docx locale : {e}, bascule vers Google Docs API...")
        build_with_google_docs()
